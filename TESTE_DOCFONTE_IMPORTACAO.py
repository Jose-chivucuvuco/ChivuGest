"""Smoke test for the SIGFE DocFonte importer.
Run: python TESTE_DOCFONTE_IMPORTACAO.py <ficheiro.xlsx>
Requires openpyxl.
"""
import sys, re, unicodedata
from collections import Counter
from openpyxl import load_workbook

def normalize_key(k):
    s = unicodedata.normalize('NFKD', str(k or '')).encode('ascii','ignore').decode().lower().strip()
    s = re.sub(r'[^a-z0-9]+','_',s).strip('_')
    return {
        'beneficiario':'fornecedor','n_os':'n_os','no_os':'n_os','situacao_os':'situacao',
        'data_emissao_os':'data_os','data_confirmacao_pagamento':'data_confirmacao_pagamento',
        'valor_total_mn':'valor_total_mn','valor_os_mn':'valor_os_mn',
        'finalidade_da_os':'finalidade_os','no_contrato':'numero_contrato','no_bancario':'numero_bancario'
    }.get(s,s)

def rv(row,*keys):
    row={normalize_key(k):v for k,v in row.items()}
    for k in keys:
        k=normalize_key(k)
        if k in row and row[k] not in (None,''): return row[k]
    return ''

p=sys.argv[1] if len(sys.argv)>1 else 'Fic_Doc_Fonte_79249_20261006092429(1).xlsx'
wb=load_workbook(p,read_only=True,data_only=True); ws=wb.active
it=ws.iter_rows(min_row=1,max_row=5001,max_col=100,values_only=True)
headers=[normalize_key(x) for x in next(it)]
rows=[dict(zip(headers,r)) for r in it if any(x not in (None,'') for x in r)]
with_os=[r for r in rows if str(rv(r,'numero_os','n_os')).strip()]
print('Total de linhas:', len(rows))
print('Linhas com Nº OS:', len(with_os))
print('Linhas sem Nº OS:', len(rows)-len(with_os))
print('Fornecedor/Beneficiário reconhecido:', sum(bool(rv(r,'beneficiario','fornecedor')) for r in with_os))
print('Situação OS reconhecida:', sum(bool(rv(r,'situacao_os','situacao')) for r in with_os))
print('Valor MN utilizável (Total MN ou OS MN):', sum(bool(rv(r,'valor_total_mn','valor_os_mn')) for r in with_os))
print('Finalidade reconhecida:', sum(bool(rv(r,'finalidade_da_os','finalidade_os')) for r in with_os))
print('Primeira OS:', rv(with_os[0],'numero_os','n_os') if with_os else '')
print('Primeiro fornecedor:', rv(with_os[0],'beneficiario','fornecedor') if with_os else '')
print('Primeiro valor MN:', rv(with_os[0],'valor_total_mn') if with_os else '')
