import os, csv, io, re, json, hashlib, secrets, unicodedata, xml.etree.ElementTree as ET
from datetime import datetime, date, timedelta
from functools import wraps
from decimal import Decimal, InvalidOperation
from flask import Flask, render_template, request, redirect, url_for, session, flash, Response
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from sqlalchemy import func, inspect, text
from sqlalchemy.exc import IntegrityError

try:
    import fitz
except ImportError:
    fitz = None
try:
    import pytesseract
except ImportError:
    pytesseract = None
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None
from PIL import Image

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", secrets.token_hex(32))
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024

database_url = os.environ.get("DATABASE_URL", "sqlite:///chivugest.db")
if database_url.startswith("postgres://"):
    database_url = database_url.replace("postgres://", "postgresql+psycopg://", 1)
elif database_url.startswith("postgresql://"):
    database_url = database_url.replace("postgresql://", "postgresql+psycopg://", 1)
app.config["SQLALCHEMY_DATABASE_URI"] = database_url
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
db = SQLAlchemy(app)

# -----------------------------------------------------------------------------
# Existing entities kept for backwards compatibility with the first ChivuGest
# release. The new procurement/supplier layer is additive and does not delete
# existing data.
# -----------------------------------------------------------------------------
class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(150), nullable=False)
    username = db.Column(db.String(80), unique=True, nullable=False, index=True)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(30), nullable=False, default="user")
    active = db.Column(db.Boolean, nullable=False, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class Client(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False, index=True)
    nif = db.Column(db.String(30))
    phone = db.Column(db.String(50))
    email = db.Column(db.String(150))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class Invoice(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    number = db.Column(db.String(80), nullable=False, index=True)
    client_id = db.Column(db.Integer, db.ForeignKey("client.id"), nullable=False)
    date = db.Column(db.Date, nullable=False)
    due_date = db.Column(db.Date, nullable=False)
    amount = db.Column(db.Numeric(18, 2), nullable=False)
    paid = db.Column(db.Numeric(18, 2), nullable=False, default=0)
    status = db.Column(db.String(20), nullable=False, default="Pendente")
    client = db.relationship("Client", backref="invoices")

class Payment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    receipt = db.Column(db.String(80), nullable=False, index=True)
    client_id = db.Column(db.Integer, db.ForeignKey("client.id"), nullable=False)
    invoice_id = db.Column(db.Integer, db.ForeignKey("invoice.id"))
    date = db.Column(db.Date, nullable=False)
    method = db.Column(db.String(50), nullable=False)
    amount = db.Column(db.Numeric(18, 2), nullable=False)
    client = db.relationship("Client", backref="payments")
    invoice = db.relationship("Invoice", backref="payments")

# -----------------------------------------------------------------------------
# Supplier / procurement / contract management
# -----------------------------------------------------------------------------
class Supplier(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(220), nullable=False, index=True)
    nif = db.Column(db.String(40), index=True)
    phone = db.Column(db.String(60))
    email = db.Column(db.String(180))
    address = db.Column(db.String(300))
    category = db.Column(db.String(120))
    contracting_type = db.Column(db.String(80), nullable=False)
    portal_status = db.Column(db.String(80), default="Não verificado")
    certification_status = db.Column(db.String(80), default="Não informado")
    tax_clearance_expiry = db.Column(db.Date)
    social_security_expiry = db.Column(db.Date)
    professional_license_expiry = db.Column(db.Date)
    blocked = db.Column(db.Boolean, nullable=False, default=False)
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class ProcurementProcedure(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    code = db.Column(db.String(100), nullable=False, unique=True)
    object = db.Column(db.Text, nullable=False)
    contract_category = db.Column(db.String(60), nullable=False, default="Bens")
    procedure_type = db.Column(db.String(80), nullable=False)
    estimated_value = db.Column(db.Numeric(18,2), nullable=False, default=0)
    budget_year = db.Column(db.Integer, default=lambda: date.today().year)
    budgeted = db.Column(db.Boolean, default=False)
    cabimentado = db.Column(db.Boolean, default=False)
    cabimentacao_ref = db.Column(db.String(120))
    decision_date = db.Column(db.Date)
    invitation_date = db.Column(db.Date)
    proposal_deadline = db.Column(db.Date)
    adjudication_date = db.Column(db.Date)
    portal_registered = db.Column(db.Boolean, default=False)
    legal_basis = db.Column(db.String(250))
    justification = db.Column(db.Text)
    status = db.Column(db.String(50), default="Em preparação")
    supplier_id = db.Column(db.Integer, db.ForeignKey("supplier.id"))
    created_by = db.Column(db.Integer, db.ForeignKey("user.id"))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    supplier = db.relationship("Supplier")
    creator = db.relationship("User")

class Contract(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    number = db.Column(db.String(100), nullable=False, unique=True)
    supplier_id = db.Column(db.Integer, db.ForeignKey("supplier.id"), nullable=False)
    procedure_id = db.Column(db.Integer, db.ForeignKey("procurement_procedure.id"))
    object = db.Column(db.Text, nullable=False)
    contract_type = db.Column(db.String(80), nullable=False)
    procedure_type = db.Column(db.String(80), nullable=False)
    start_date = db.Column(db.Date, nullable=False)
    end_date = db.Column(db.Date, nullable=False)
    original_value = db.Column(db.Numeric(18,2), nullable=False, default=0)
    current_value = db.Column(db.Numeric(18,2), nullable=False, default=0)
    renewal_allowed = db.Column(db.Boolean, default=False)
    renewal_count = db.Column(db.Integer, default=0)
    cabimentado = db.Column(db.Boolean, default=False)
    cabimentacao_ref = db.Column(db.String(120))
    tribunal_review_required = db.Column(db.Boolean, default=False)
    tribunal_review_status = db.Column(db.String(60), default="Não aplicável")
    guarantee_required = db.Column(db.Boolean, default=False)
    guarantee_value = db.Column(db.Numeric(18,2), default=0)
    advance_percent = db.Column(db.Numeric(7,2), default=0)
    amendments_percent = db.Column(db.Numeric(7,2), default=0)
    status = db.Column(db.String(40), default="Em vigor")
    document_ref = db.Column(db.String(300))
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    supplier = db.relationship("Supplier", backref="contracts")
    procedure = db.relationship("ProcurementProcedure", backref="contracts")

class SupplierInvoice(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    number = db.Column(db.String(100), nullable=False, index=True)
    supplier_id = db.Column(db.Integer, db.ForeignKey("supplier.id"), nullable=False)
    contract_id = db.Column(db.Integer, db.ForeignKey("contract.id"))
    issue_date = db.Column(db.Date, nullable=False)
    due_date = db.Column(db.Date)
    subtotal = db.Column(db.Numeric(18,2), default=0)
    vat = db.Column(db.Numeric(18,2), default=0)
    total = db.Column(db.Numeric(18,2), nullable=False, default=0)
    paid = db.Column(db.Numeric(18,2), nullable=False, default=0)
    currency = db.Column(db.String(10), default="AOA")
    status = db.Column(db.String(30), default="Pendente")
    source_filename = db.Column(db.String(255))
    source_hash = db.Column(db.String(64), unique=True)
    raw_text = db.Column(db.Text)
    extracted_data = db.Column(db.Text)
    import_confidence = db.Column(db.Integer, default=0)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    supplier = db.relationship("Supplier", backref="invoices")
    contract = db.relationship("Contract", backref="invoices")

class SupplierPayment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    receipt = db.Column(db.String(100), nullable=False, index=True)
    supplier_id = db.Column(db.Integer, db.ForeignKey("supplier.id"), nullable=False)
    invoice_id = db.Column(db.Integer, db.ForeignKey("supplier_invoice.id"))
    contract_id = db.Column(db.Integer, db.ForeignKey("contract.id"))
    date = db.Column(db.Date, nullable=False)
    method = db.Column(db.String(60), nullable=False)
    amount = db.Column(db.Numeric(18,2), nullable=False)
    reference = db.Column(db.String(150))
    notes = db.Column(db.Text)
    supplier = db.relationship("Supplier", backref="payments")
    invoice = db.relationship("SupplierInvoice", backref="payments")
    contract = db.relationship("Contract", backref="payments")

class LegalRule(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    code = db.Column(db.String(100), unique=True, nullable=False)
    title = db.Column(db.String(220), nullable=False)
    source = db.Column(db.String(250), nullable=False)
    article = db.Column(db.String(80))
    summary = db.Column(db.Text, nullable=False)
    severity = db.Column(db.String(20), default="ALERTA")
    active = db.Column(db.Boolean, default=True)
    parameter = db.Column(db.String(80))
    value = db.Column(db.Numeric(18,2))
    source_url = db.Column(db.String(500))
    version = db.Column(db.String(50), default="2026")

class ComplianceAlert(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    alert_type = db.Column(db.String(80), nullable=False)
    severity = db.Column(db.String(20), nullable=False, default="ALERTA")
    title = db.Column(db.String(250), nullable=False)
    message = db.Column(db.Text, nullable=False)
    legal_basis = db.Column(db.String(250))
    supplier_id = db.Column(db.Integer, db.ForeignKey("supplier.id"))
    contract_id = db.Column(db.Integer, db.ForeignKey("contract.id"))
    procedure_id = db.Column(db.Integer, db.ForeignKey("procurement_procedure.id"))
    resolved = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    supplier = db.relationship("Supplier")
    contract = db.relationship("Contract")
    procedure = db.relationship("ProcurementProcedure")

class LegalDocument(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(250), nullable=False)
    description = db.Column(db.Text)
    source_url = db.Column(db.String(500), nullable=False)
    current_version = db.Column(db.String(60), nullable=False)
    active = db.Column(db.Boolean, default=True)

# -----------------------------------------------------------------------------
# Helpers
# -----------------------------------------------------------------------------
PROCEDURES = [
    "Concurso Público", "Concurso Limitado por Prévia Qualificação",
    "Concurso Limitado por Convite", "Contratação Simplificada",
    "Procedimento Dinâmico Electrónico", "Contratação Emergencial"
]
CONTRACT_CATEGORIES = ["Empreitada", "Bens", "Serviços", "Concessão", "Locação"]
CONTRACTING_TYPES = PROCEDURES + ["Outro / Regime especial"]


def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if not session.get("uid"):
            return redirect(url_for("login"))
        u = db.session.get(User, session["uid"])
        if not u or not u.active:
            session.clear(); return redirect(url_for("login"))
        return f(*a, **k)
    return w


def admin_required(f):
    @wraps(f)
    def w(*a, **k):
        if session.get("role") != "admin":
            flash("Acesso reservado ao administrador.")
            return redirect(url_for("dashboard"))
        return f(*a, **k)
    return w

@app.context_processor
def globals_processor():
    return {"today": date.today(), "procedures": PROCEDURES, "contract_categories": CONTRACT_CATEGORIES,
            "contracting_types": CONTRACTING_TYPES}


def parse_date(v, default="__TODAY__"):
    if v is None or str(v).strip() == "":
        return date.today() if default == "__TODAY__" else default
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d"):
        try: return datetime.strptime(s, fmt).date()
        except ValueError: pass
    return default or date.today()


def money(v):
    try: return float(v or 0)
    except Exception: return 0.0


def num(v):
    if v is None or str(v).strip() == "": return 0.0
    s = str(v).strip().upper().replace("AOA", "").replace("AKZ", "").replace("KZ", "").replace("$", "").replace("€", "")
    s = re.sub(r"[^0-9,.-]", "", s)
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."): s = s.replace(".", "").replace(",", ".")
        else: s = s.replace(",", "")
    elif "," in s:
        # 1,234.56 was handled above; here comma is normally decimal in Angola/Portugal.
        s = s.replace(",", ".")
    try: return float(s)
    except (ValueError, TypeError): return 0.0


def normalize_key(k):
    if k is None: return ""
    s = unicodedata.normalize("NFKD", str(k)).encode("ascii", "ignore").decode().lower().strip()
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_")
    aliases = {
        "factura":"fatura", "n_fatura":"numero", "numero_fatura":"numero", "n_fatura":"numero",
        "invoice_number":"numero", "invoice_no":"numero", "documento":"numero", "doc":"numero",
        "fornecedor":"fornecedor", "supplier":"fornecedor", "nome_fornecedor":"fornecedor", "emitente":"fornecedor",
        "nuit":"nif", "vat_number":"nif", "tax_id":"nif", "tin":"nif",
        "data_emissao":"data", "issue_date":"data", "invoice_date":"data", "data_fatura":"data",
        "due_date":"vencimento", "data_vencimento":"vencimento", "data_limite":"vencimento",
        "valor_total":"total", "total_fatura":"total", "valor":"total", "amount":"total", "grand_total":"total",
        "subtotal":"subtotal", "base_tributavel":"subtotal", "iva":"iva", "vat":"iva", "imposto":"iva",
        "moeda":"currency", "meio_pagamento":"metodo", "payment_method":"metodo", "metodo_pagamento":"metodo",
        "iban":"iban", "referencia":"referencia", "reference":"referencia", "descricao":"descricao", "description":"descricao"
    }
    return aliases.get(s, s)


def row_value(row, *keys):
    normalized = {normalize_key(k): v for k, v in row.items()}
    for key in keys:
        nk = normalize_key(key)
        if nk in normalized and normalized[nk] not in (None, ""):
            return normalized[nk]
    return ""


def normalize_text(text_value):
    return re.sub(r"[ \t]+", " ", text_value.replace("\xa0", " "))


def regex_first(text_value, patterns):
    for p in patterns:
        m = re.search(p, text_value, re.I | re.M)
        if m:
            return m.group(1).strip(" :#-\t")
    return ""


def parse_invoice_text(text_value):
    text_value = normalize_text(text_value)
    number = regex_first(text_value, [
        r"(?:fatura|factura|invoice|documento)\s*(?:n[ºo°]?|no|numero|número|number)?\s*[:#-]?\s*([A-Z]{0,5}[A-Z0-9./_-]{2,})",
        r"\b((?:FT|FA|FR|FAC|INV)[\s./_-]*[A-Z0-9_-]{2,})\b"
    ])
    nif = regex_first(text_value, [r"(?:NIF|NUIT|N\.I\.F\.)\s*[:#-]?\s*([0-9]{8,15})"])
    supplier = regex_first(text_value, [
        r"(?:fornecedor|supplier|emitente|vendedor)\s*[:#-]\s*([^\n\r]{3,120})",
        r"(?:nome|name)\s*[:#-]\s*([^\n\r]{3,120})"
    ])
    dates = re.findall(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}[/-]\d{1,2}[/-]\d{1,2})\b", text_value)
    total = regex_first(text_value, [
        r"(?:total a pagar|total da fatura|total factura|total fatura|grand total|valor total|total)\s*[:=\-]?\s*(?:AOA|KZ|AKZ)?\s*([0-9][0-9 .,'-]{1,})",
        r"(?:valor a pagar|montante)\s*[:=\-]?\s*(?:AOA|KZ|AKZ)?\s*([0-9][0-9 .,'-]{1,})"
    ])
    subtotal = regex_first(text_value, [r"(?:subtotal|total liquido|base tributavel)\s*[:=\-]?\s*(?:AOA|KZ|AKZ)?\s*([0-9][0-9 .,'-]{1,})"])
    vat = regex_first(text_value, [r"(?:IVA|VAT)\s*(?:\([0-9]+(?:[.,][0-9]+)?%\))?\s*[:=\-]?\s*(?:AOA|KZ|AKZ)?\s*([0-9][0-9 .,'-]{1,})"])
    iban = regex_first(text_value, [r"\b(\w{2}\d{2}[A-Z0-9 ]{10,34})\b"])
    currency = "AOA" if re.search(r"\b(?:KZ|AOA|AKZ|KWANZA|KWANZAS)\b", text_value, re.I) else ""
    data = {
        "number": number, "supplier": supplier, "nif": nif,
        "date": dates[0] if dates else "", "due_date": dates[1] if len(dates) > 1 else "",
        "subtotal": num(subtotal), "vat": num(vat), "total": num(total), "currency": currency,
        "iban": iban.replace(" ", ""), "text": text_value
    }
    if not data["total"] and data["subtotal"]:
        data["total"] = data["subtotal"] + data["vat"]
    score = 0
    score += 25 if data["number"] else 0
    score += 20 if data["supplier"] else 0
    score += 15 if data["nif"] else 0
    score += 15 if data["date"] else 0
    score += 25 if data["total"] else 0
    data["confidence"] = score
    return data


def extract_invoice_document(file_storage):
    filename=(file_storage.filename or "").lower()
    if filename.endswith((".xml", ".json")):
        raw=file_storage.read(); digest=hashlib.sha256(raw).hexdigest()
        if filename.endswith(".json"):
            obj=json.loads(raw.decode("utf-8-sig"))
            if isinstance(obj, list): obj=obj[0] if obj else {}
            flat=[]
            def walk(x,prefix=""):
                if isinstance(x,dict):
                    for k,v in x.items(): walk(v, f"{prefix}.{k}" if prefix else str(k))
                elif isinstance(x,list):
                    for i,v in enumerate(x): walk(v, f"{prefix}.{i}")
                else: flat.append((prefix,x))
            walk(obj)
            text_json="\n".join(f"{k}: {v}" for k,v in flat)
        else:
            root=ET.fromstring(raw)
            pairs=[]
            for el in root.iter():
                if el.text and el.text.strip(): pairs.append((el.tag.split('}')[-1], el.text.strip()))
            text_json="\n".join(f"{k}: {v}" for k,v in pairs)
        parsed=parse_invoice_text(text_json); parsed.update({"hash":digest,"filename":file_storage.filename})
        # Structured documents often have clearer field names than OCR.
        try:
            if filename.endswith(".json"):
                obj=json.loads(raw.decode("utf-8-sig")); flat_text=json.dumps(obj,ensure_ascii=False)
            else: flat_text=text_json
            parsed["supplier"] = parsed["supplier"] or regex_first(flat_text,[r"(?:SupplierName|Supplier|Fornecedor|Seller)\s*[:=]\s*[\"']?([^,\n\"']+)"])
            parsed["number"] = parsed["number"] or regex_first(flat_text,[r"(?:InvoiceNumber|InvoiceNo|Numero|Fatura|Factura)\s*[:=]\s*[\"']?([A-Z0-9./_-]+)"])
        except Exception:
            pass
        return parsed
    return extract_pdf_or_image(file_storage)


def extract_pdf_or_image(file_storage):
    filename = (file_storage.filename or "").lower()
    raw = file_storage.read()
    digest = hashlib.sha256(raw).hexdigest()
    texts = []
    if filename.endswith(".pdf"):
        if fitz is None: raise ValueError("PyMuPDF não instalado.")
        doc = fitz.open(stream=raw, filetype="pdf")
        for page in doc:
            txt = page.get_text("text") or ""
            if len(re.sub(r"\s+", "", txt)) < 40:
                if pytesseract is None: raise ValueError("OCR não instalado.")
                pix = page.get_pixmap(matrix=fitz.Matrix(2.4, 2.4), alpha=False)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                try: txt = pytesseract.image_to_string(img, lang="por+eng")
                except Exception: txt = pytesseract.image_to_string(img, lang="eng")
            texts.append(txt)
        doc.close()
    elif filename.endswith((".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff")):
        if pytesseract is None: raise ValueError("OCR não instalado.")
        img = Image.open(io.BytesIO(raw)).convert("RGB")
        try: texts.append(pytesseract.image_to_string(img, lang="por+eng"))
        except Exception: texts.append(pytesseract.image_to_string(img, lang="eng"))
    else:
        raise ValueError("Formato não suportado para OCR. Use PDF, PNG, JPG, JPEG, WEBP, TIF ou TIFF.")
    parsed = parse_invoice_text("\n".join(texts))
    parsed["hash"] = digest
    parsed["filename"] = file_storage.filename
    return parsed


def uploaded_rows(f):
    name = (f.filename or "").lower()
    if name.endswith(".xlsx"):
        if not load_workbook: raise ValueError("openpyxl não instalado.")
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True)); wb.close()
        if not data: return []
        headers = [normalize_key(x) for x in data[0]]
        return [dict(zip(headers, r)) for r in data[1:] if any(x is not None for x in r)]
    if name.endswith(".csv"):
        raw = f.read()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try: s = raw.decode(enc); break
            except UnicodeDecodeError: s = None
        if s is None: raise ValueError("Não foi possível ler o CSV.")
        sample = s[:4096]
        try: dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";" if ";" in sample else ","
        return [{normalize_key(k): v for k, v in row.items()} for row in csv.DictReader(io.StringIO(s), dialect=dialect)]
    raise ValueError("Para tabelas use CSV/XLSX. Para documentos use PDF/imagem.")

# -----------------------------------------------------------------------------
# Legal rules and compliance engine. This is a rules assistant, not a legal
# opinion. Thresholds are stored in the DB so administrators can update them
# when the annual OGE/Annex I changes.
# -----------------------------------------------------------------------------
LEGAL_DOCS = [
    ("Lei n.º 41/20 — Lei dos Contratos Públicos", "Regime jurídico da formação e execução dos contratos públicos.", "https://lex.ao/docs/assembleia-nacional/2020/lei-n-o-41-20-de-23-de-dezembro/", "41/20"),
    ("Decreto Presidencial n.º 74/26 — REOGE 2026", "Regras de execução do OGE 2026, incluindo execução de contratos públicos.", "https://www.angolex.com/paginas/decreto-presidencial/regras-de-execucao-do-orcamento-geral-do-estado-para-2026a-74a-26a.html", "74/26"),
    ("Decreto Presidencial n.º 77/23", "Procedimento de cobrança e destino das coimas da contratação pública.", "https://lex.ao/docs/presidente-da-republica/2023/decreto-presidencial-n-o-77-23-de-20-de-marco/", "77/23"),
    ("Decreto Presidencial n.º 250/24", "Plano Estratégico da Contratação Pública Angolana 2024–2028.", "https://lex.ao/docs/presidente-da-republica/2024/decreto-presidencial-n-o-250-24-de-13-de-novembro/", "250/24"),
]

LEGAL_SEEDS = [
    ("LCP_PROC_VALUE", "Escolha do procedimento por valor", "Lei 41/20", "Art. 22.º–25.º", "Concurso Público ou Concurso Limitado por Prévia Qualificação é a regra; Concurso Limitado por Convite e Contratação Simplificada têm limites de valor.", "CRITICO", None, None),
    ("LCP_LOTS", "Valor agregado dos lotes", "Lei 41/20", "Art. 25.º", "Quando prestações do mesmo tipo são divididas em lotes, o valor considerado para escolha do procedimento é o somatório dos valores estimados dos lotes.", "CRITICO", None, None),
    ("LCP_SIMPLIFIED", "Contratação Simplificada por valor", "Lei 41/20", "Art. 24.º/27.º–30.º", "A contratação simplificada por valor não deve exceder o Nível 1 do Anexo I; critérios materiais têm regime próprio.", "CRITICO", "nivel_1", 18000000),
    ("LCP_INVITATION", "Concurso Limitado por Convite", "Lei 41/20", "Art. 24.º", "O Concurso Limitado por Convite só permite contratos abaixo do Nível 5 do Anexo I.", "CRITICO", "nivel_5", 182000000),
    ("LCP_CONCESSION", "Concessões", "Lei 41/20", "Art. 24.º, n.º 5", "Para concessões deve ser adotado Concurso Público ou Concurso Limitado por Prévia Qualificação, independentemente do valor.", "CRITICO", None, None),
    ("LCP_BUDGET", "Decisão de contratar e verba", "Lei 41/20", "Art. 32.º", "A decisão de contratar pressupõe verba inscrita no orçamento, salvo a condição legalmente prevista.", "CRITICO", None, None),
    ("LCP_WRITTEN", "Contrato escrito", "Lei 41/20", "Art. 106.º–107.º", "Regra geral: contrato reduzido a escrito, com exceções legais por valor/natureza.", "ALERTA", None, None),
    ("LCP_BOND", "Caução", "Lei 41/20", "Art. 99.º", "A caução é obrigatória para adjudicações iguais ou superiores ao Nível 5, sem prejuízo de exigência em valor inferior.", "ALERTA", None, None),
    ("LCP_SANCTIONS", "Impedimentos e sanções", "Lei 41/20", "Art. 56.º–57.º e 429.º–438.º", "O sistema deve verificar impedimentos, empresas incumpridoras e situações sancionatórias registadas.", "CRITICO", None, None),
    ("LCP_PAC", "Plano Anual de Contratação", "Lei 41/20", "Art. 442.º", "A entidade pública deve elaborar e comunicar o Plano Anual de Contratação.", "ALERTA", None, None),
    ("REOGE_CAB", "Cabimentação prévia", "DP 74/26", "Art. 9.º–10.º", "É vedado iniciar despesas, obras, contratos ou requisições sem prévia cabimentação nos limites legais.", "CRITICO", None, None),
    ("REOGE_ADV_WORKS", "Adiantamento de empreitada", "DP 74/26", "Art. 10.º", "Adiantamento inicial de empreitada até 15%; até 30% pode ser autorizado pelo responsável das Finanças nos termos do diploma.", "CRITICO", "adiantamento_empreitada", 15),
    ("REOGE_ADV_GOODS", "Pagamento inicial de bens/serviços", "DP 74/26", "Art. 10.º", "Pagamento inicial para bens e serviços de despesas correntes pode ir até 50%, desde que devidamente fundamentado.", "CRITICO", "adiantamento_bens_servicos", 50),
    ("REOGE_AMEND", "Adendas/trabalhos a mais", "DP 74/26", "Art. 10.º", "É proibida adenda resultante de trabalhos a mais quando o valor total excede 15% do contrato inicial, sem prejuízo do regime de reequilíbrio financeiro.", "CRITICO", "adendas", 15),
    ("REOGE_RENEW", "Contratos contínuos", "DP 74/26", "Art. 10.º", "Contratos contínuos podem ser prorrogados em condições legais até ao prazo máximo de 48 meses; depois exige-se novo procedimento.", "CRITICO", "meses_continuos", 48),
    ("REOGE_HIGH_VALUE", "Diligências de beneficiário efetivo", "DP 74/26", "Art. 10.º", "Acima de Kz 500M em empreitadas e Kz 182M em bens/serviços, devem ser observadas as diligências previstas no diploma quando solicitadas pelos bancos.", "ALERTA", "alto_valor_bens_servicos", 182000000),
]


def seed_legal_data():
    for name, desc, url, version in LEGAL_DOCS:
        if not LegalDocument.query.filter_by(name=name).first():
            db.session.add(LegalDocument(name=name, description=desc, source_url=url, current_version=version))
    for code, title, source, article, summary, severity, parameter, value in LEGAL_SEEDS:
        rule = LegalRule.query.filter_by(code=code).first()
        if not rule:
            db.session.add(LegalRule(code=code, title=title, source=source, article=article,
                                     summary=summary, severity=severity, parameter=parameter, value=value,
                                     source_url=next((u for n,d,u,v in LEGAL_DOCS if n.startswith(source)), None)))
    db.session.commit()


def get_rule_value(code, default):
    r = LegalRule.query.filter_by(code=code, active=True).first()
    return float(r.value) if r and r.value is not None else default


def add_alert(alert_type, severity, title, message, legal_basis="", supplier_id=None, contract_id=None, procedure_id=None):
    existing = ComplianceAlert.query.filter_by(alert_type=alert_type, contract_id=contract_id,
                                                procedure_id=procedure_id, supplier_id=supplier_id,
                                                resolved=False).first()
    if not existing:
        db.session.add(ComplianceAlert(alert_type=alert_type, severity=severity, title=title,
                                       message=message, legal_basis=legal_basis, supplier_id=supplier_id,
                                       contract_id=contract_id, procedure_id=procedure_id))


def run_compliance_checks():
    today = date.today()
    # Supplier-document and eligibility alerts.
    for s in Supplier.query.all():
        if not s.contracting_type:
            add_alert("FORNECEDOR_SEM_TIPO", "CRITICO", "Fornecedor sem tipo de contratação",
                      f"O fornecedor {s.name} não tem tipo de contratação definido.", "Controlo interno", s.id)
        if s.blocked or str(s.portal_status).lower().startswith("bloque"):
            add_alert("FORNECEDOR_BLOQUEADO", "CRITICO", "Fornecedor bloqueado",
                      f"O fornecedor {s.name} está marcado como bloqueado. Rever antes de adjudicar, contratar ou pagar.",
                      "Lei 41/20, Art. 56.º–57.º", s.id)
        for label, expiry in (("Certidão fiscal", s.tax_clearance_expiry), ("Segurança Social", s.social_security_expiry), ("Licença profissional", s.professional_license_expiry)):
            if expiry:
                days=(expiry-today).days
                if days < 0:
                    add_alert("DOC_FORNECEDOR_EXPIRADO", "CRITICO", f"{label} expirada",
                              f"{label} do fornecedor {s.name} expirou em {expiry.strftime('%d/%m/%Y')}.", "Controlo de habilitação do fornecedor", s.id)
                elif days <= 30:
                    add_alert("DOC_FORNECEDOR_A_EXPIRAR", "ALERTA", f"{label} a expirar",
                              f"{label} do fornecedor {s.name} expira em {days} dia(s).", "Controlo de habilitação do fornecedor", s.id)
    # Contracts and expiration alerts.
    for c in Contract.query.all():
        days = (c.end_date - today).days
        if c.status == "Em vigor" and days < 0:
            add_alert("CONTRATO_EXPIRADO", "CRITICO", "Contrato expirado",
                      f"O contrato {c.number} terminou em {c.end_date.strftime('%d/%m/%Y')} e continua marcado como em vigor.",
                      "Gestão contratual / DP 74/26 — prorrogações devem observar limites legais.", c.supplier_id, c.id)
        elif c.status == "Em vigor" and days <= 7:
            add_alert("CONTRATO_7_DIAS", "CRITICO", "Contrato vence em breve",
                      f"O contrato {c.number} vence em {days} dia(s). Inicie a regularização/novo procedimento.",
                      "Gestão contratual", c.supplier_id, c.id)
        elif c.status == "Em vigor" and days <= 30:
            add_alert("CONTRATO_30_DIAS", "ALERTA", "Contrato vence em 30 dias",
                      f"O contrato {c.number} vence em {days} dia(s).", "Gestão contratual", c.supplier_id, c.id)
        elif c.status == "Em vigor" and days <= 60:
            add_alert("CONTRATO_60_DIAS", "INFO", "Contrato aproxima-se do vencimento",
                      f"O contrato {c.number} vence em {days} dia(s).", "Gestão contratual", c.supplier_id, c.id)
        if c.amendments_percent and float(c.amendments_percent) > get_rule_value("REOGE_AMEND", 15):
            add_alert("ADENDA_EXCESSIVA", "CRITICO", "Adendas acima do limite parametrizado",
                      f"O contrato {c.number} acumula {c.amendments_percent}% em adendas, acima do limite de {get_rule_value('REOGE_AMEND',15):g}%.",
                      "DP 74/26, Art. 10.º", c.supplier_id, c.id)
        if c.advance_percent:
            limit = get_rule_value("REOGE_ADV_WORKS", 15) if c.contract_type == "Empreitada" else get_rule_value("REOGE_ADV_GOODS", 50)
            if float(c.advance_percent) > limit:
                add_alert("ADIANTAMENTO_EXCESSIVO", "CRITICO", "Adiantamento acima do limite parametrizado",
                          f"O contrato {c.number} tem adiantamento de {c.advance_percent}%, acima do limite base de {limit:g}%.",
                          "DP 74/26, Art. 10.º", c.supplier_id, c.id)
        if not c.cabimentado:
            add_alert("SEM_CABIMENTACAO", "CRITICO", "Contrato sem cabimentação registada",
                      f"O contrato {c.number} está marcado sem cabimentação. Não deve avançar a execução sem a condição orçamental aplicável.",
                      "Lei 41/20, Art. 32.º; DP 74/26, Art. 10.º", c.supplier_id, c.id)
        if c.contract_type in ("Bens", "Serviços") and (c.end_date - c.start_date).days > 48*31:
            add_alert("PRAZO_CONTINUO", "CRITICO", "Prazo superior ao limite de 48 meses",
                      f"O contrato {c.number} ultrapassa aproximadamente 48 meses de vigência.",
                      "DP 74/26, Art. 10.º", c.supplier_id, c.id)
        if float(c.current_value or 0) >= get_rule_value("REOGE_HIGH_VALUE", 182000000) and c.contract_type in ("Bens", "Serviços"):
            add_alert("ALTO_VALOR", "ALERTA", "Contrato de alto valor",
                      f"O contrato {c.number} ultrapassa Kz {get_rule_value('REOGE_HIGH_VALUE',182000000):,.0f}. Rever diligências de beneficiário efetivo e documentação financeira quando legalmente solicitadas.",
                      "DP 74/26, Art. 10.º", c.supplier_id, c.id)
    # Procedure checks.
    level1 = get_rule_value("LCP_SIMPLIFIED", 18000000)
    level5 = get_rule_value("LCP_INVITATION", 182000000)
    for p in ProcurementProcedure.query.all():
        v = float(p.estimated_value or 0)
        if not p.budgeted or not p.cabimentado:
            add_alert("PROCEDIMENTO_SEM_ORCAMENTO", "CRITICO", "Procedimento sem cobertura/cabimentação registada",
                      f"O procedimento {p.code} não tem orçamento/cabimentação marcados como regulares.",
                      "Lei 41/20, Art. 32.º; DP 74/26, Art. 10.º", p.supplier_id, None, p.id)
        if p.contract_category == "Concessão" and p.procedure_type not in ("Concurso Público", "Concurso Limitado por Prévia Qualificação"):
            add_alert("PROCEDIMENTO_INADEQUADO", "CRITICO", "Procedimento incompatível com concessão",
                      f"O procedimento {p.code} usa {p.procedure_type} para concessão.", "Lei 41/20, Art. 24.º, n.º 5", p.supplier_id, None, p.id)
        if p.procedure_type == "Contratação Simplificada" and v > level1:
            add_alert("SIMPLIFICADA_ACIMA_LIMITE", "CRITICO", "Contratação Simplificada acima do limite",
                      f"O procedimento {p.code} tem valor estimado de Kz {v:,.2f}, acima do Nível 1 parametrizado (Kz {level1:,.2f}). Só deve prosseguir se existir fundamento material legal aplicável.",
                      "Lei 41/20, Art. 24.º e 27.º–30.º", p.supplier_id, None, p.id)
        if p.procedure_type == "Concurso Limitado por Convite" and v >= level5:
            add_alert("CONVITE_ACIMA_LIMITE", "CRITICO", "Concurso Limitado por Convite fora do limite",
                      f"O procedimento {p.code} tem valor estimado de Kz {v:,.2f}, não inferior ao limite parametrizado do Nível 5 (Kz {level5:,.2f}).", "Lei 41/20, Art. 24.º", p.supplier_id, None, p.id)
        if p.procedure_type in ("Contratação Simplificada", "Contratação Emergencial") and not p.justification:
            add_alert("SEM_FUNDAMENTACAO", "ALERTA", "Procedimento excepcional sem fundamentação registada",
                      f"O procedimento {p.code} exige revisão da fundamentação e do critério material/emergencial, conforme aplicável.", "Lei 41/20, Art. 26.º–31.º", p.supplier_id, None, p.id)
        if not p.portal_registered:
            add_alert("PORTAL_NAO_REGISTADO", "ALERTA", "Registo no Portal não confirmado",
                      f"O procedimento {p.code} está marcado sem registo confirmado no Portal da Contratação Pública.", "Lei 41/20, Art. 12.º e regras do procedimento", p.supplier_id, None, p.id)
    db.session.commit()

# -----------------------------------------------------------------------------
# Auth
# -----------------------------------------------------------------------------
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = User.query.filter_by(username=request.form.get("username", "").strip()).first()
        if u and u.active and check_password_hash(u.password_hash, request.form.get("password", "")):
            session.update(uid=u.id, name=u.name, role=u.role)
            return redirect(url_for("dashboard"))
        flash("Utilizador ou palavra-passe inválidos.")
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear(); return redirect(url_for("login"))

# -----------------------------------------------------------------------------
# Dashboard
# -----------------------------------------------------------------------------
@app.route("/")
@login_required
def dashboard():
    total_invoices = db.session.query(func.coalesce(func.sum(SupplierInvoice.total), 0)).scalar() or 0
    total_paid = db.session.query(func.coalesce(func.sum(SupplierPayment.amount), 0)).scalar() or 0
    payable = float(total_invoices) - float(total_paid)
    overdue = db.session.query(func.coalesce(func.sum(SupplierInvoice.total - SupplierInvoice.paid), 0)).filter(
        SupplierInvoice.due_date < date.today(), SupplierInvoice.total > SupplierInvoice.paid).scalar() or 0
    contracts_active = Contract.query.filter_by(status="Em vigor").count()
    expiring = Contract.query.filter(Contract.status == "Em vigor", Contract.end_date <= date.today()+timedelta(days=60), Contract.end_date >= date.today()).count()
    critical_alerts = ComplianceAlert.query.filter_by(resolved=False, severity="CRITICO").count()
    alert_count = ComplianceAlert.query.filter_by(resolved=False).count()
    suppliers = Supplier.query.count()
    recent_payments = SupplierPayment.query.order_by(SupplierPayment.id.desc()).limit(8).all()
    recent_contracts = Contract.query.order_by(Contract.end_date).limit(6).all()
    monthly = db.session.query(func.extract("year", SupplierPayment.date).label("y"), func.extract("month", SupplierPayment.date).label("m"), func.sum(SupplierPayment.amount).label("v")).group_by("y","m").order_by("y","m").all()
    months = [{"label": f"{int(r.y):04d}-{int(r.m):02d}", "value": float(r.v or 0)} for r in monthly]
    maxv = max([m["value"] for m in months], default=0)
    for m in months: m["height"] = 20 + (m["value"]/(maxv or 1))*180
    run_compliance_checks()
    return render_template("dashboard.html", total_invoices=total_invoices, total_paid=total_paid, payable=payable,
                           overdue=overdue, contracts_active=contracts_active, expiring=expiring,
                           critical_alerts=critical_alerts, alert_count=alert_count, suppliers=suppliers,
                           recent_payments=recent_payments, recent_contracts=recent_contracts, months=months)

# -----------------------------------------------------------------------------
# Suppliers
# -----------------------------------------------------------------------------
@app.route("/suppliers", methods=["GET", "POST"])
@login_required
def suppliers():
    if request.method == "POST":
        try:
            s = Supplier(name=request.form["name"].strip(), nif=request.form.get("nif"), phone=request.form.get("phone"),
                         email=request.form.get("email"), address=request.form.get("address"), category=request.form.get("category"),
                         contracting_type=request.form["contracting_type"], portal_status=request.form.get("portal_status","Não verificado"),
                         certification_status=request.form.get("certification_status","Não informado"),
                         tax_clearance_expiry=parse_date(request.form.get("tax_clearance_expiry"), None),
                         social_security_expiry=parse_date(request.form.get("social_security_expiry"), None),
                         professional_license_expiry=parse_date(request.form.get("professional_license_expiry"), None),
                         blocked=bool(request.form.get("blocked")), notes=request.form.get("notes"))
            db.session.add(s); db.session.commit(); flash("Fornecedor criado. O tipo de contratação é obrigatório para acompanhamento.")
        except Exception as e:
            db.session.rollback(); flash("Não foi possível criar o fornecedor: "+str(e))
    return render_template("suppliers.html", rows=Supplier.query.order_by(Supplier.name).all())

# -----------------------------------------------------------------------------
# Supplier invoices and payments
# -----------------------------------------------------------------------------
@app.route("/invoices", methods=["GET", "POST"])
@login_required
def supplier_invoices():
    suppliers_list = Supplier.query.order_by(Supplier.name).all()
    contracts = Contract.query.filter_by(status="Em vigor").order_by(Contract.number).all()
    if request.method == "POST":
        try:
            supplier = db.session.get(Supplier, int(request.form["supplier_id"]))
            total = num(request.form["total"])
            paid = num(request.form.get("paid"))
            inv = SupplierInvoice(number=request.form["number"].strip(), supplier_id=supplier.id,
                contract_id=int(request.form["contract_id"]) if request.form.get("contract_id") else None,
                issue_date=parse_date(request.form.get("issue_date")), due_date=parse_date(request.form.get("due_date"), None),
                subtotal=num(request.form.get("subtotal")), vat=num(request.form.get("vat")), total=total, paid=paid,
                currency=request.form.get("currency","AOA"), status="Paga" if paid>=total and total>0 else ("Parcial" if paid>0 else "Pendente"),
                source_filename=request.form.get("source_filename"))
            db.session.add(inv); db.session.commit(); flash("Fatura de fornecedor registada.")
        except Exception as e:
            db.session.rollback(); flash("Erro ao registar fatura: "+str(e))
    return render_template("supplier_invoices.html", rows=SupplierInvoice.query.order_by(SupplierInvoice.id.desc()).all(), suppliers=suppliers_list, contracts=contracts)

@app.route("/payments", methods=["GET", "POST"])
@login_required
def supplier_payments():
    suppliers_list = Supplier.query.order_by(Supplier.name).all()
    invoices = SupplierInvoice.query.order_by(SupplierInvoice.id.desc()).all()
    contracts = Contract.query.filter_by(status="Em vigor").all()
    if request.method == "POST":
        try:
            amount = num(request.form["amount"])
            inv = db.session.get(SupplierInvoice, int(request.form["invoice_id"])) if request.form.get("invoice_id") else None
            supplier_id = int(request.form["supplier_id"])
            if inv:
                supplier_id = inv.supplier_id
                inv.paid = money(inv.paid) + amount
                inv.status = "Paga" if inv.paid >= inv.total else "Parcial"
            p = SupplierPayment(receipt=request.form["receipt"], supplier_id=supplier_id,
                                invoice_id=inv.id if inv else None,
                                contract_id=int(request.form["contract_id"]) if request.form.get("contract_id") else None,
                                date=parse_date(request.form.get("date")), method=request.form["method"], amount=amount,
                                reference=request.form.get("reference"), notes=request.form.get("notes"))
            db.session.add(p); db.session.commit(); flash("Pagamento registado.")
        except Exception as e:
            db.session.rollback(); flash("Erro no pagamento: "+str(e))
    return render_template("supplier_payments.html", rows=SupplierPayment.query.order_by(SupplierPayment.id.desc()).all(), suppliers=suppliers_list, invoices=invoices, contracts=contracts)

@app.route("/current")
@login_required
def current_account():
    supplier_id = request.args.get("supplier_id", type=int)
    suppliers_list = Supplier.query.order_by(Supplier.name).all()
    inv_q = SupplierInvoice.query.order_by(SupplierInvoice.issue_date, SupplierInvoice.id)
    pay_q = SupplierPayment.query.order_by(SupplierPayment.date, SupplierPayment.id)
    if supplier_id:
        inv_q = inv_q.filter_by(supplier_id=supplier_id); pay_q = pay_q.filter_by(supplier_id=supplier_id)
    entries=[]
    for i in inv_q.all(): entries.append((i.issue_date, i.number, "Fatura", money(i.total), 0, i.supplier.name))
    for p in pay_q.all(): entries.append((p.date, p.receipt, "Pagamento", 0, money(p.amount), p.supplier.name))
    entries.sort(key=lambda x:(x[0],x[1] or "")); bal=0; rows=[]
    for e in entries:
        bal += e[3]-e[4]; rows.append((e,bal))
    return render_template("current.html", rows=rows, suppliers=suppliers_list, supplier_id=supplier_id)

# -----------------------------------------------------------------------------
# Contracts and procurement
# -----------------------------------------------------------------------------
@app.route("/contracts", methods=["GET", "POST"])
@login_required
def contracts():
    suppliers_list=Supplier.query.order_by(Supplier.name).all(); procedures=ProcurementProcedure.query.order_by(ProcurementProcedure.code).all()
    if request.method == "POST":
        try:
            c=Contract(number=request.form["number"], supplier_id=int(request.form["supplier_id"]),
                procedure_id=int(request.form["procedure_id"]) if request.form.get("procedure_id") else None,
                object=request.form["object"], contract_type=request.form["contract_type"], procedure_type=request.form["procedure_type"],
                start_date=parse_date(request.form["start_date"]), end_date=parse_date(request.form["end_date"]),
                original_value=num(request.form["original_value"]), current_value=num(request.form.get("current_value") or request.form["original_value"]),
                renewal_allowed=bool(request.form.get("renewal_allowed")), renewal_count=int(request.form.get("renewal_count") or 0),
                cabimentado=bool(request.form.get("cabimentado")), cabimentacao_ref=request.form.get("cabimentacao_ref"),
                tribunal_review_required=bool(request.form.get("tribunal_review_required")), tribunal_review_status=request.form.get("tribunal_review_status","Não aplicável"),
                guarantee_required=bool(request.form.get("guarantee_required")), guarantee_value=num(request.form.get("guarantee_value")),
                advance_percent=num(request.form.get("advance_percent")), amendments_percent=num(request.form.get("amendments_percent")),
                status=request.form.get("status","Em vigor"), document_ref=request.form.get("document_ref"), notes=request.form.get("notes"))
            db.session.add(c); db.session.commit(); run_compliance_checks(); flash("Contrato registado e submetido ao motor de conformidade.")
        except Exception as e: db.session.rollback(); flash("Erro no contrato: "+str(e))
    return render_template("contracts.html", rows=Contract.query.order_by(Contract.end_date).all(), suppliers=suppliers_list, procedures=procedures)

@app.route("/procurement", methods=["GET", "POST"])
@login_required
def procurement():
    suppliers_list=Supplier.query.order_by(Supplier.name).all()
    if request.method == "POST":
        try:
            p=ProcurementProcedure(code=request.form["code"], object=request.form["object"], contract_category=request.form["contract_category"],
                procedure_type=request.form["procedure_type"], estimated_value=num(request.form["estimated_value"]), budget_year=int(request.form.get("budget_year") or date.today().year),
                budgeted=bool(request.form.get("budgeted")), cabimentado=bool(request.form.get("cabimentado")), cabimentacao_ref=request.form.get("cabimentacao_ref"),
                decision_date=parse_date(request.form.get("decision_date"), None), invitation_date=parse_date(request.form.get("invitation_date"), None),
                proposal_deadline=parse_date(request.form.get("proposal_deadline"), None), adjudication_date=parse_date(request.form.get("adjudication_date"), None),
                portal_registered=bool(request.form.get("portal_registered")), legal_basis=request.form.get("legal_basis"), justification=request.form.get("justification"),
                status=request.form.get("status","Em preparação"), supplier_id=int(request.form["supplier_id"]) if request.form.get("supplier_id") else None, created_by=session["uid"])
            db.session.add(p); db.session.commit(); run_compliance_checks(); flash("Procedimento registado e analisado.")
        except Exception as e: db.session.rollback(); flash("Erro no procedimento: "+str(e))
    return render_template("procurement.html", rows=ProcurementProcedure.query.order_by(ProcurementProcedure.id.desc()).all(), suppliers=suppliers_list)

@app.route("/alerts")
@login_required
def alerts():
    run_compliance_checks()
    return render_template("alerts.html", rows=ComplianceAlert.query.filter_by(resolved=False).order_by(ComplianceAlert.created_at.desc()).all())

@app.route("/alerts/resolve/<int:aid>")
@login_required
def resolve_alert(aid):
    a=db.session.get(ComplianceAlert, aid)
    if a: a.resolved=True; db.session.commit()
    return redirect(url_for("alerts"))

# -----------------------------------------------------------------------------
# Reports
# -----------------------------------------------------------------------------
@app.route("/reports")
@login_required
def reports():
    total=money(db.session.query(func.coalesce(func.sum(SupplierInvoice.total),0)).scalar())
    paid=money(db.session.query(func.coalesce(func.sum(SupplierPayment.amount),0)).scalar())
    return render_template("reports.html", total=total, paid=paid, payable=total-paid,
                           suppliers=Supplier.query.count(), contracts=Contract.query.count(), procedures=ProcurementProcedure.query.count(),
                           alerts=ComplianceAlert.query.filter_by(resolved=False).count(), invoices=SupplierInvoice.query.count())

@app.route("/reports/export")
@login_required
def report_export():
    out=io.StringIO(); w=csv.writer(out)
    w.writerow(["Fornecedor","NIF","Contrato","Fatura","Data","Vencimento","Total","Pago","Saldo","Estado"])
    for i in SupplierInvoice.query.order_by(SupplierInvoice.issue_date).all():
        w.writerow([i.supplier.name,i.supplier.nif or "",i.contract.number if i.contract else "",i.number,i.issue_date.isoformat(),i.due_date.isoformat() if i.due_date else "",money(i.total),money(i.paid),money(i.total)-money(i.paid),i.status])
    return Response("\ufeff"+out.getvalue(),mimetype="text/csv",headers={"Content-Disposition":"attachment; filename=chivugest_relatorio_fornecedores.csv"})

# -----------------------------------------------------------------------------
# Robust import centre
# -----------------------------------------------------------------------------
@app.route("/import", methods=["GET", "POST"])
@login_required
def import_center():
    if request.method == "POST":
        kind=request.form.get("kind")
        f=request.files.get("file")
        if not f or not f.filename: flash("Selecione um ficheiro."); return redirect(url_for("import_center"))
        try:
            if kind == "supplier_invoices_table":
                rows=uploaded_rows(f); n=0
                for r in rows:
                    supplier_name=str(row_value(r,"fornecedor","supplier","emitente")).strip()
                    number=str(row_value(r,"numero","fatura","factura","invoice_number")).strip()
                    total=num(row_value(r,"total","valor","amount"))
                    if not supplier_name or not number or total <= 0: continue
                    s=Supplier.query.filter_by(name=supplier_name).first()
                    if not s:
                        s=Supplier(name=supplier_name,nif=str(row_value(r,"nif","nuit")),contracting_type="Outro / Regime especial")
                        db.session.add(s); db.session.flush()
                    inv=SupplierInvoice(number=number,supplier_id=s.id,issue_date=parse_date(row_value(r,"data","issue_date")),
                        due_date=parse_date(row_value(r,"vencimento","due_date"),None),subtotal=num(row_value(r,"subtotal")),vat=num(row_value(r,"iva","vat")),
                        total=total,paid=num(row_value(r,"pago","paid")),currency=str(row_value(r,"currency","moeda")) or "AOA",
                        source_filename=f.filename)
                    inv.status="Paga" if inv.paid>=inv.total else ("Parcial" if inv.paid>0 else "Pendente")
                    db.session.add(inv); n+=1
                db.session.commit(); flash(f"Importação concluída: {n} fatura(s).")
            elif kind == "pdf_invoice":
                parsed=extract_invoice_document(f)
                existing=SupplierInvoice.query.filter_by(source_hash=parsed["hash"]).first()
                if existing: raise ValueError("Este documento já foi importado anteriormente.")
                session["invoice_import"] = parsed
                return render_template("invoice_review.html", data=parsed)
            else:
                raise ValueError("Tipo de importação desconhecido.")
        except Exception as e:
            db.session.rollback(); flash("Erro na importação: "+str(e))
    return render_template("import_center.html")

@app.route("/import/invoice-confirm", methods=["POST"])
@login_required
def invoice_confirm():
    data=session.pop("invoice_import", {})
    try:
        supplier_name=request.form["supplier"]
        s=Supplier.query.filter_by(name=supplier_name).first()
        if not s:
            s=Supplier(name=supplier_name,nif=request.form.get("nif"),contracting_type=request.form.get("contracting_type") or "Outro / Regime especial")
            db.session.add(s); db.session.flush()
        inv=SupplierInvoice(number=request.form["number"], supplier_id=s.id, issue_date=parse_date(request.form.get("date")),
            due_date=parse_date(request.form.get("due_date"),None), subtotal=num(request.form.get("subtotal")), vat=num(request.form.get("vat")),
            total=num(request.form.get("total")), paid=0, currency=request.form.get("currency") or "AOA", source_filename=data.get("filename"),
            source_hash=data.get("hash"), raw_text="", extracted_data=json.dumps(data,ensure_ascii=False), import_confidence=int(data.get("confidence",0)))
        inv.status="Pendente"; db.session.add(inv); db.session.commit(); flash("Fatura importada após revisão.")
    except Exception as e: db.session.rollback(); flash("Erro ao confirmar a fatura: "+str(e))
    return redirect(url_for("supplier_invoices"))

# -----------------------------------------------------------------------------
# Settings / legal library / users
# -----------------------------------------------------------------------------
@app.route("/settings", methods=["GET", "POST"])
@login_required
@admin_required
def settings():
    if request.method == "POST":
        for rule in LegalRule.query.all():
            if rule.parameter and rule.parameter in request.form:
                try: rule.value=Decimal(request.form[rule.parameter])
                except InvalidOperation: pass
        db.session.commit(); flash("Parâmetros legais atualizados.")
    return render_template("settings.html", rules=LegalRule.query.order_by(LegalRule.id).all(), documents=LegalDocument.query.order_by(LegalDocument.id).all())

@app.route("/legal")
@login_required
def legal_library():
    return render_template("legal.html", documents=LegalDocument.query.order_by(LegalDocument.id).all(), rules=LegalRule.query.filter_by(active=True).order_by(LegalRule.id).all())

@app.route("/users", methods=["GET", "POST"])
@login_required
@admin_required
def users():
    if request.method=="POST":
        try:
            db.session.add(User(name=request.form["name"], username=request.form["username"], password_hash=generate_password_hash(request.form["password"]), role=request.form["role"], active=True))
            db.session.commit(); flash("Utilizador criado.")
        except IntegrityError: db.session.rollback(); flash("Nome de utilizador já existe.")
    return render_template("users.html", rows=User.query.order_by(User.name).all())

@app.route("/users/toggle/<int:uid>")
@login_required
@admin_required
def toggle_user(uid):
    u=db.session.get(User,uid)
    if u and u.id != session["uid"]: u.active=not u.active; db.session.commit()
    return redirect(url_for("users"))

# -----------------------------------------------------------------------------
# Legacy customer module kept available for backwards compatibility.
# -----------------------------------------------------------------------------
@app.route("/clients", methods=["GET", "POST"])
@login_required
def clients():
    if request.method=="POST":
        db.session.add(Client(name=request.form["name"],nif=request.form.get("nif"),phone=request.form.get("phone"),email=request.form.get("email"))); db.session.commit(); flash("Cliente criado.")
    return render_template("clients.html", rows=Client.query.order_by(Client.name).all())

@app.route("/legacy-invoices", methods=["GET", "POST"])
@login_required
def legacy_invoices():
    clients=Client.query.order_by(Client.name).all()
    if request.method=="POST":
        db.session.add(Invoice(number=request.form["number"],client_id=int(request.form["client_id"]),date=parse_date(request.form["date"]),due_date=parse_date(request.form["due_date"]),amount=num(request.form["amount"]),paid=0,status="Pendente")); db.session.commit(); flash("Fatura de cliente registada.")
    return render_template("invoices.html", rows=Invoice.query.order_by(Invoice.id.desc()).all(), clients=clients)

@app.route("/export")
@login_required
def export():
    out=io.StringIO(); w=csv.writer(out); w.writerow(["Fornecedor","NIF","Fatura","Data","Vencimento","Total","Pago","Saldo","Estado"])
    for i in SupplierInvoice.query.order_by(SupplierInvoice.issue_date).all(): w.writerow([i.supplier.name,i.supplier.nif or "",i.number,i.issue_date.isoformat(),i.due_date.isoformat() if i.due_date else "",money(i.total),money(i.paid),money(i.total)-money(i.paid),i.status])
    return Response("\ufeff"+out.getvalue(),mimetype="text/csv",headers={"Content-Disposition":"attachment; filename=chivugest_faturas_fornecedores.csv"})

@app.route("/health")
def health(): return "OK", 200

# -----------------------------------------------------------------------------
# Database initialisation. Existing tables are preserved. New tables are added.
# -----------------------------------------------------------------------------
def init_db():
    with app.app_context():
        db.create_all()
        if not User.query.filter_by(username="admin").first():
            db.session.add(User(name="Administrador",username="admin",password_hash=generate_password_hash("admin123"),role="admin"))
            db.session.commit()
        seed_legal_data()
        # Make sure legacy installations that had no new tables are initialized.
        run_compliance_checks()

init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT",5000)), debug=False)
